"""social.workflow — the media generation orchestrator.

  generate_media    read Draft rows -> generate the missing ones -> upload
                    route-by-type -> write the Drive link + cost back into the sheet

(upload / download / snapshot of the calendar itself live in social.sheet_ops.)

generate_media takes ``mode`` = dry-run | mock | live:
  dry-run  plan only. No Drive, no API, nothing written.
  mock     real pipeline with placeholder files. Uploads and write-back are
           routed to a SAFE mock destination (a mock root, or a '_mock rehearsal'
           subfolder) and a *.mock.xlsx copy — production assets are never touched.
  live     the real thing.

Only Status=Draft rows are ever generated, and a row is skipped if its asset is
already on Drive (deleting the Drive file is how you request a regeneration).
"""

from __future__ import annotations

import sys
from pathlib import Path

from . import rules
from .calendar import Calendar
from ..core import config, media
from ..core.drive import FOLDER_MIME, file_id_from_link

MOCK_SUBFOLDER = "_mock rehearsal"
FAILED_TEXT = "Failed"  # written into the asset-link cell when a row can't be produced


def _uses_selected(job) -> bool:
    """True when this row should be filled from its Selected Asset Link instead of
    generating. Applies to single-image and single-video rows; a lone link can't
    populate a multi-slide carousel, so carousels always generate."""
    return bool(job.selected_link) and job.plan.kind in ("image", "video")


_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp", ".gif")
_VIDEO_EXTS = (".mp4", ".mov", ".m4v", ".webm")


def _drive_target_from_link(drive, link: str, ext: str) -> str | None:
    """The concrete Drive FILE id for a Selected Asset Drive link. A file link resolves to
    itself; a FOLDER link (…/drive/folders/<id>) resolves to the single file inside that
    matches the target type (image for .png, video for .mp4). Raises a clear error if the
    folder is empty of matches or holds several candidates — link the specific file then."""
    fid = file_id_from_link(link)
    if not fid or "/folders/" not in link:
        return fid  # a direct file link (or unparseable -> caller raises)
    exts = _VIDEO_EXTS if ext.lower() in _VIDEO_EXTS else _IMAGE_EXTS
    kindword = "video" if exts is _VIDEO_EXTS else "image"
    files = [c for c in drive.list_children(fid) if c.get("mimeType") != FOLDER_MIME]
    matches = [c for c in files if c["name"].lower().endswith(exts)]
    if len(matches) == 1:
        return matches[0]["id"]
    if not matches:
        raise FileNotFoundError(
            f"the linked folder has no {kindword} file to copy — "
            "link the specific file, not the folder")
    raise ValueError(
        f"the linked folder has {len(matches)} {kindword} files — "
        "link the specific file, not the folder")


def _fetch_selected(drive, link: str, out_path: Path) -> None:
    """Place the Selected Asset (a Drive file/folder share link, a plain http(s) URL, or a
    local path) at ``out_path``. Raises if the source can't be reached/read so the caller
    can record the failure and move on."""
    import shutil

    link = link.strip()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if "drive.google" in link or "docs.google" in link:
        fid = _drive_target_from_link(drive, link, out_path.suffix)
        if not fid:
            raise ValueError(f"could not parse a Drive file id from link: {link}")
        drive.download_file(fid, out_path)
    elif link.lower().startswith(("http://", "https://")):
        import urllib.request
        urllib.request.urlretrieve(link, out_path)  # noqa: S310 (trusted, human-entered)
    else:
        src = Path(link)
        if not src.is_file():
            raise FileNotFoundError(f"selected asset not found: {link}")
        shutil.copyfile(src, out_path)


def _use_selected_asset(drive, job, out_dir: Path, emit) -> dict:
    """A run_batch-shaped receipt built by copying the row's Selected Asset into the
    location the generated image would have occupied — no image model is called. If the
    source can't be fetched, the asset is recorded as an error (and the row is later
    marked Failed like any other generation failure) and we keep going."""
    receipt = {"batch_id": job.row_id, "outputs": [], "errors": [], "hints": []}
    for asset in job.assets:
        target_dir = media.asset_target_dir(out_dir, asset)
        ext = ".mp4" if asset.get("type") == "video" else ".png"  # match the model's output name
        for n in media.revision_numbers(asset):
            out_path = target_dir / f"{asset['id']}_v{n}{ext}"
            try:
                _fetch_selected(drive, job.selected_link, out_path)
                emit(f"  select -> {out_path.name}  (copied from Selected Asset Link)")
                receipt["outputs"].append({"file": str(out_path), "model": "selected-asset",
                                           "est_cost_usd": 0.0})
            except Exception as e:
                short, _ = media.friendly_error(e)
                receipt["errors"].append({"id": asset.get("id"), "error": str(e),
                                          "reason": f"selected asset unavailable: {short}"})
                emit(f"  FAILED -> {asset.get('id')}: selected asset unavailable ({short})",
                     err=True)
    return receipt


def _stderr_emit(msg: str, *, err: bool = False) -> None:
    print(msg, file=sys.stderr)


def _drive_client(interactive: bool = False):
    from ..core.drive import DriveClient
    return DriveClient(config.credentials_path(), config.token_path(),
                       allow_interactive=interactive)


def _resolve_base_folder(drive, calendar_id: str, mode: str) -> tuple[str, str]:
    """Return (base_folder_id, folder_name). For live this is the calendar's folder
    under the Social Calendar root; for mock it's a sandbox that mirrors it."""
    root = rules.social_calendar_root_id()
    if not root:
        raise RuntimeError("SOCIAL_CALENDAR_ROOT_ID is not set — point it at the "
                           "'Social Calendar' Drive folder id.")
    folder = rules.calendar_folder(calendar_id)
    if mode == "mock":
        mock_root = rules.social_calendar_mock_root_id()
        if mock_root:
            return drive.ensure_path(mock_root, [folder]), folder
        folder_id = drive.find_or_create_folder(folder, root)
        return drive.ensure_path(folder_id, [MOCK_SUBFOLDER]), folder
    return drive.find_or_create_folder(folder, root), folder


def _kind_parent(drive, base_folder: str, job) -> str:
    """The route-by-type folder to check for an existing asset, and (for image/video)
    to upload into. Carousels check/inhabit the 03_Carousels parent; the per-set group
    subfolder is created only when actually generating — not during the existence check."""
    if job.plan.kind == "carousel":
        return drive.ensure_path(base_folder, list(rules.SUBFOLDER_CAROUSELS))
    if job.plan.recorded:  # Wiah's own clip -> 01_Wiah Videos (not the AI video folder)
        return drive.ensure_path(base_folder, list(rules.SUBFOLDER_WIAH_VIDEOS))
    if job.plan.kind == "video":
        return drive.ensure_path(base_folder, list(rules.SUBFOLDER_VIDEO))
    return drive.ensure_path(base_folder, list(rules.SUBFOLDER_IMAGES))


def _already_on_drive(drive, parent: str, job) -> str | None:
    """Existing Drive link if this row's asset is already there, else None. Keyed on
    the stable Row ID prefix (``{row_id}_``), NOT the hook-derived slug — so editing a
    caption never orphans an already-generated file. A carousel counts as present when
    a group subfolder named ``{row_id}_...`` exists and holds at least one file."""
    if job.plan.kind == "carousel":
        for child in drive.list_children(parent):
            if (child.get("mimeType") == FOLDER_MIME
                    and child["name"].startswith(f"{job.row_id}_")):
                has_file = any(f.get("mimeType") != FOLDER_MIME
                               for f in drive.list_children(child["id"]))
                if has_file:
                    return drive.get_link(child["id"])
        return None
    hits = drive.find_by_prefix(f"{job.row_id}_", parent)
    return hits[0].get("webViewLink") if hits else None


def _drive_md5(drive, file_id: str) -> str | None:
    try:
        return drive.get_file(file_id).get("md5Checksum")
    except Exception:
        return None


def _selected_source_md5(drive, link: str, ext: str) -> str | None:
    """The MD5 of the Selected Asset's content: a Drive file's checksum (resolving a folder
    link to its one matching file, like the copy does) or a local file's hash. None for a
    remote URL, a Google-native file, an ambiguous folder, or anything we can't checksum —
    callers treat None as 'can't verify' and re-copy to be safe."""
    link = (link or "").strip()
    if "drive.google" in link or "docs.google" in link:
        try:
            fid = _drive_target_from_link(drive, link, ext)
        except Exception:
            return None  # empty/ambiguous folder -> re-copy; the copy will error clearly
        return _drive_md5(drive, fid) if fid else None
    if link.lower().startswith(("http://", "https://")):
        return None  # don't refetch a remote URL just to compare
    try:
        import hashlib
        p = Path(link)
        if p.is_file():
            return hashlib.md5(p.read_bytes()).hexdigest()
    except Exception:
        pass
    return None


def _selected_asset_changed(drive, parent: str, job) -> bool:
    """For a Selected-Asset row that already has a copy on Drive: True if the Selected
    Asset now points at DIFFERENT content than that copy (=> re-copy), False if it's the
    same (=> skip). Compared by MD5 (Drive's checksum equals a local hash for identical
    bytes). If either side can't be checksummed we return True, so an intended change is
    never silently skipped."""
    ext = ".mp4" if job.plan.kind == "video" else ".png"
    sel = _selected_source_md5(drive, job.selected_link, ext)
    if sel is None:
        return True
    hits = drive.find_by_prefix(f"{job.row_id}_", parent)
    if not hits:
        return True
    return _drive_md5(drive, hits[0]["id"]) != sel


# --- writeback target -------------------------------------------------------
class _SheetWriter:
    """Records machine-owned cell edits and flushes them to the living Google Sheet in
    place via the Sheets API — same write_result/write_note interface as Calendar, so the
    generate loop is writeback-agnostic (mock writes to the openpyxl Calendar instead)."""

    def __init__(self, sheets, spreadsheet_id: str, tab: str, cal):
        self.sheets, self.sid, self.tab, self.cal = sheets, spreadsheet_id, tab, cal
        self.updates: list[tuple[str, object]] = []

    def _a1(self, field: str, row: int) -> str | None:
        from openpyxl.utils import get_column_letter
        c = self.cal.cols.get(field)
        return f"'{self.tab}'!{get_column_letter(c)}{row}" if c else None

    def write_result(self, row_index, link=None, cost=None, model=None, overwrite=True):
        for field, val in (("asset_link", link), ("est_cost", cost), ("ai_model", model)):
            if val is None:
                continue
            # non-live rehearsal: only fill a blank cell, never clobber a real value
            if not overwrite and not Calendar.is_blank(self.cal._get(row_index, field)):
                continue
            a1 = self._a1(field, row_index)
            if a1:
                self.updates.append((a1, val))

    def write_note(self, row_index, text):
        a1 = self._a1("notes", row_index)
        if a1:
            self.updates.append(
                (a1, Calendar.merged_note(self.cal._get(row_index, "notes"), text)))

    def flush(self) -> int:
        return self.sheets.batch_update(self.sid, self.updates).get(
            "totalUpdatedCells", len(self.updates))


# --- operations ------------------------------------------------------------
def generate_media(calendar_id: str, mode: str = "dry-run", *,
                   only: str | None = None,
                   image_model: str | None = None, video_model: str | None = None,
                   video_duration: int | None = None, emit=None) -> dict:
    """Generate the missing AI visuals for the living calendar's Draft rows, upload them,
    and write the Drive link / cost / model / notes back INTO THE LIVING GOOGLE SHEET in
    place (live) — no download/re-upload, so concurrent human edits aren't clobbered.
    dry-run plans + costs only; mock rehearses to a safe Drive destination and a local
    *.mock.xlsx (the live sheet is untouched).

    image_model / video_model / video_duration override per-row settings for this run
    only; on success the model actually used is written into the sheet's AI Model cell."""
    import io
    emit = emit or _stderr_emit
    if mode not in ("dry-run", "mock", "live"):
        raise ValueError(f"mode must be dry-run|mock|live, got {mode!r}")

    # The calendar is the living Google Sheet, exported to .xlsx bytes for reading.
    from . import sheet_ops
    drive = _drive_client(interactive=False)
    docs = sheet_ops._docs_folder(drive, calendar_id)
    live = sheet_ops.find_live_sheet(drive, docs, calendar_id)
    if not live:
        raise FileNotFoundError(
            f"no live sheet '{sheet_ops.live_sheet_name(calendar_id)}' on Drive; "
            "run `social upload <id> <version>` first.")
    sid = live["id"]
    cal = Calendar(io.BytesIO(drive.export_as_xlsx(sid)))
    jobs = cal.read_jobs()

    def _wanted(job) -> bool:
        if not job.in_scope:
            return False
        if only == "image":
            return job.plan.kind in ("image", "carousel")
        if only == "video":
            return job.plan.kind == "video"
        return True

    in_scope = [j for j in jobs if _wanted(j)]
    # Overrides: explicit arg (flag/param) beats the env default, which beats the sheet.
    video_model = video_model or config.video_model_override()
    image_model = image_model or config.image_model_override()
    video_duration = video_duration if video_duration is not None else config.video_duration_override()
    for job in in_scope:
        for a in job.assets:
            if a["type"] == "video":
                if video_model:
                    a["model"] = video_model
                if video_duration:
                    a["duration_seconds"] = video_duration
            elif a["type"] == "image" and image_model:
                a["model"] = image_model
    out_dir = config.generated_dir() / calendar_id
    result = {"calendar_id": calendar_id, "mode": mode, "rows_total": len(jobs),
              "in_scope": len(in_scope), "generated": 0, "skipped_existing": 0,
              "failed": 0, "estimated_cost_usd": 0.0, "rows": [], "hints": [],
              "sheet_link": live.get("webViewLink")}

    # --- dry-run: plan + cost only; the live sheet is not modified ------------
    if mode == "dry-run":
        total = 0.0
        for j in in_scope:
            if _uses_selected(j):
                # Copied from the Selected Asset Link — no model call, no cost.
                result["rows"].append({"row_id": j.row_id, "kind": j.plan.kind,
                                       "aspect_ratio": j.plan.aspect_ratio,
                                       "assets": len(j.assets), "action": "would-copy-selected",
                                       "cost_usd": 0.0})
                continue
            rec = media.run_batch(j.assets, defaults=media.DEFAULTS, out_dir=out_dir,
                                  mode="dry-run", emit=emit, batch_id=j.row_id)
            total += rec["estimated_cost_usd"]
            result["rows"].append({"row_id": j.row_id, "kind": j.plan.kind,
                                   "aspect_ratio": j.plan.aspect_ratio,
                                   "assets": len(j.assets), "action": "would-generate",
                                   "cost_usd": round(rec["estimated_cost_usd"], 4)})
        result["estimated_cost_usd"] = round(total, 2)
        result["note"] = ("dry-run: worst-case plan + cost only. Nothing is generated and "
                          "the live sheet is not modified.")
        return result

    # --- mock / live: Drive-backed -------------------------------------------
    base_folder, folder = _resolve_base_folder(drive, calendar_id, mode)
    result["folder"] = folder
    image_client = video_client = types = None
    if mode == "live":
        # Init only the client(s) the in-scope assets actually need: images -> OpenAI
        # (gpt-image-2), video -> google-genai (Veo). An image-only run never needs a
        # Gemini key, and vice-versa.
        # Rows filled from a Selected Asset Link never call the image model, so a run
        # made up entirely of those doesn't need an OpenAI key/client.
        if any(a["type"] == "image" for j in in_scope if not _uses_selected(j)
               for a in j.assets):
            image_client = media.init_image_client()
        if any(a["type"] == "video" for j in in_scope if not _uses_selected(j)
               for a in j.assets):
            video_client, types = media.init_video_client()
        from ..core.sheets import SheetsClient
        writer = _SheetWriter(SheetsClient(config.credentials_path(), config.token_path()),
                              sid, cal.ws.title, cal)
    else:  # mock: openpyxl copy saved to a local *.mock.xlsx; live sheet untouched
        writer = cal

    # Only a live run overwrites an existing cell; a mock rehearsal fills blanks only,
    # so its estimates/placeholders never clobber real values carried from the sheet.
    overwrite = mode == "live"

    hints: set[str] = set()
    for job in in_scope:
        parent = _kind_parent(drive, base_folder, job)
        existing = _already_on_drive(drive, parent, job)
        # A Selected-Asset row is only "already done" if the on-Drive copy still matches
        # the Selected Asset Link; if the link now points at a different file, re-copy.
        if existing and _uses_selected(job) and _selected_asset_changed(drive, parent, job):
            emit(f"  update -> {job.row_id} (Selected Asset changed; re-copying)")
            existing = None
        if existing:
            result["skipped_existing"] += 1
            # A copied (Selected Asset) row — incl. recorded-Wiah clips — has no
            # generation cost; only truly generated assets carry an estimate.
            est = 0.0 if _uses_selected(job) else media.estimate_cost(job.assets)
            result["rows"].append({"row_id": job.row_id, "kind": job.plan.kind,
                                   "action": "skipped-existing", "link": existing,
                                   "cost_usd": est})
            writer.write_result(job.row_index, link=existing, cost=est, overwrite=overwrite)
            writer.write_note(job.row_index, "")  # clear any stale failure note
            continue

        if _uses_selected(job):
            # Copy the human-picked asset into the target location instead of generating.
            rec = _use_selected_asset(drive, job, out_dir, emit)
        else:
            rec = media.run_batch(job.assets, defaults=media.DEFAULTS, out_dir=out_dir,
                                  mode=mode, emit=emit, batch_id=job.row_id,
                                  image_client=image_client, video_client=video_client,
                                  types=types)
        hints.update(rec.get("hints", []))
        outputs = [o for o in rec["outputs"] if not o.get("dry_run")]
        if rec["errors"] or not outputs:
            result["failed"] += 1
            err = rec["errors"][0]["reason"] if rec["errors"] else "no output produced"
            writer.write_result(job.row_index, link=FAILED_TEXT, overwrite=overwrite)
            writer.write_note(job.row_index, f"generate failed: {err}")
            result["rows"].append({"row_id": job.row_id, "kind": job.plan.kind,
                                   "action": "failed", "error": err})
            continue

        dest = (drive.ensure_path(parent, [job.group])
                if job.plan.kind == "carousel" else parent)
        try:
            link = _upload_row(drive, dest, job, outputs)
        except Exception as e:
            short, hint = media.friendly_error(e)
            if hint:
                hints.add(hint)
            result["failed"] += 1
            writer.write_result(job.row_index, link=FAILED_TEXT, overwrite=overwrite)
            writer.write_note(job.row_index, f"upload failed: {short}")
            result["rows"].append({"row_id": job.row_id, "kind": job.plan.kind,
                                   "action": "upload-failed", "error": short})
            continue

        cost = round(sum(o.get("est_cost_usd", 0) for o in outputs), 4)
        writer.write_result(job.row_index, link=link, cost=cost, overwrite=overwrite,
                            model=outputs[0].get("model"))  # actual model (reflects override)
        writer.write_note(job.row_index, "")
        result["generated"] += 1
        result["estimated_cost_usd"] = round(result["estimated_cost_usd"] + cost, 4)
        result["rows"].append({"row_id": job.row_id, "kind": job.plan.kind,
                               "action": "generated", "link": link, "cost_usd": cost})

    if mode == "live":
        result["cells_written"] = writer.flush()
        emit(f"updated {result['cells_written']} cells in the live sheet in place")
    else:  # mock
        dest_path = rules.calendar_dir() / f"{sheet_ops.live_sheet_name(calendar_id)}.mock.xlsx"
        cal.save(dest_path)
        result["writeback_file"] = str(dest_path)
    result["hints"] = sorted(hints)
    emit(f"\nDone [{mode}]. generated={result['generated']} "
         f"skipped-existing={result['skipped_existing']} failed={result['failed']} "
         f"est.cost~${result['estimated_cost_usd']}.")
    return result


def _upload_row(drive, dest_folder: str, job, outputs: list[dict]) -> str | None:
    """Upload a row's generated files and return the link to write into the sheet:
    the group-folder link for a carousel, else the file link."""
    file_link = None
    for o in outputs:
        up = drive.upload(Path(o["file"]), dest_folder)
        o["drive_id"], o["drive_link"] = up["id"], up["link"]
        drive.make_shareable(up["id"])
        file_link = file_link or up["link"]
    if job.plan.kind == "carousel":
        # one folder link per carousel row (Cowork prompt step 5)
        return drive.make_shareable(dest_folder)
    return file_link


# upload / download / snapshot of the calendar live in social.sheet_ops (they operate
# on the living Google Sheet, not local .xlsx files).
